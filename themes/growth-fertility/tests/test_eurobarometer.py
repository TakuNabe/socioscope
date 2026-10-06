"""eurobarometer.py: Standard Eurobarometer "expectations for the next twelve months" (pure).

Workbook fixtures are built in-test: a small xlsx (openpyxl) mirroring the VOL A sheet layout
observed in STD93–STD105, a raw-BIFF xls (conftest) mirroring STD91/92, and zip wrappers.
"""

import io
import json
import zipfile
from collections.abc import Callable

import openpyxl
import pytest

from theme_growth_fertility import eurobarometer as eb

QUESTION_FR = "QA2.1. Quelles sont vos attentes pour les 12 prochains mois ..."
QUESTION_EN = (
    "QA2.1. What are your expectations for the next twelve months : will the next twelve "
    "months be better, worse or the same, when it comes to...? "
)
CODES = [
    "UE27\nEU27", "BE", "BG", "CZ", "DK", "D-W", "DE", "D-E", "EE", "IE", "EL", "ES", "FR", "HR",
    "IT", "CY", "LV", "LT", "LU", "HU", "MT", "NL", "AT", "PL", "PT", "RO", "SI", "SK", "FI", "SE",
    "TR", "MK", "ME", "RS", "AL", "MD", "UK", "BA", "XK", "CY(TCC)", "GE",
]  # fmt: skip


def item_sheet(
    item_en: str,
    *,
    better: list[object],
    worse: list[object],
    same: list[object],
    dk: list[object],
    same_label: str = "The same",
    dk_label: str = "Don't know",
    header_first: str = "<<Back to content",
    question: str = QUESTION_EN,
) -> list[list[object]]:
    n = len(CODES)
    rows: list[list[object]] = [
        ["Eurobarometer - 101.3"],
        ["VOL A weighted", "Terrain/Fieldwork :  02/04 - 09/05/2024"],
        [QUESTION_FR, question],
        ["Votre vie en général ", item_en],
        ["", "Base: Ensemble", "Base: All respondents"],
        [""],
        [],
        [],
        [header_first, *CODES],
        ["", "Total", *([1000] * n)],
        [],
        ["Meilleurs", *([100] * n)],
        ["Better", *better],
        ["Moins bons", *([100] * n)],
        ["Worse", *worse],
        ["Sans changement", *([100] * n)],
        [same_label, *same],
        ["Ne sait pas", *([100] * n)],
        [dk_label, *dk],
    ]
    return rows


def fill(value: object, **by_code: object) -> list[object]:
    return [by_code.get(c, value) for c in CODES]


def xlsx_from_sheets(sheets: dict[str, list[list[object]]]) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r, row in enumerate(rows, start=1):
            for c, v in enumerate(row, start=1):
                if v is not None:
                    ws.cell(row=r, column=c, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def zip_of(name: str, data: bytes, extra: dict[str, bytes] | None = None) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr(name, data)
        for k, v in (extra or {}).items():
            z.writestr(k, v)
    return buf.getvalue()


def std101() -> eb.Wave:
    return next(w for w in eb.WAVES if w.code == "STD101")


def workbook() -> bytes:
    return xlsx_from_sheets(
        {
            "Content": [["Project:", "Eurobarometer"], ["Wave:", "101.3"]],
            "QA1_1": [["QA1.1 Something else entirely"], [], ["Your life in general"]],
            "QA2_1": item_sheet(
                "Your life in general",
                better=fill(0.26, BE=0.3, FI=0.35, DE=0.17, **{"D-W": 0.16, "D-E": 0.19}),
                worse=fill(0.17, BE=0.1, FI=0.12),
                same=fill(0.55, BE=0.6, FI=0.5),
                dk=fill(0.02, BE="-", FI=0.03),
            ),
            "QA2_2": item_sheet(
                "The situation in (OUR COUNTRY) in general",
                better=fill(0.2),
                worse=fill(0.3),
                same=fill(0.4),
                dk=fill(0.1),
            ),  # fmt: skip
            "QA2_3": item_sheet(
                "The economic situation in (OUR COUNTRY)",
                better=fill(0.21),
                worse=fill(0.31),
                same=fill(0.41),
                dk=fill(0.07),
            ),  # fmt: skip
            "QA2_4": item_sheet(
                "The financial situation of your household",
                better=fill(0.22),
                worse=fill(0.32),
                same=fill(0.42),
                dk=fill(0.04),
                same_label="Same",
            ),  # fmt: skip
            "QA2_5": item_sheet(
                "The employment situation in (OUR COUNTRY)",
                better=fill(0.23),
                worse=fill(0.33),
                same=fill(0.43),
                dk=fill(0.01),
                dk_label="DK",
            ),  # fmt: skip
        }
    )


# ---------------------------------------------------------------- waves / URLs


def test_waves_are_fixed_std91_to_std105_with_unique_seasons() -> None:
    codes = [w.code for w in eb.WAVES]
    assert codes == [f"STD{n}" for n in range(91, 106)]
    assert std101().dataset_id == "s3216_101_3_std101_eng"
    assert std101().jsonld_url == (
        "https://data.europa.eu/api/hub/repo/datasets/s3216_101_3_std101_eng.jsonld"
    )
    assert std101().meta_raw_name == "eb_STD101_meta.json"
    assert (std101().fieldwork_year, std101().fieldwork_half) == (2024, 1)
    assert std101().fieldwork_start == "2024-04"
    seasons = [(w.fieldwork_year, w.fieldwork_half) for w in eb.WAVES]
    assert len(set(seasons)) == len(seasons) and seasons == sorted(seasons)
    std94 = next(w for w in eb.WAVES if w.code == "STD94")  # Winter 2020-2021, fieldwork Feb 2021
    assert (std94.fieldwork_year, std94.fieldwork_half, std94.fieldwork_start) == (
        2020,
        2,
        "2021-02",
    )
    assert eb.SOURCE == "eurobarometer_std"
    assert "CC BY 4.0" in eb.LICENSE and "2011/833/EU" in eb.LICENSE


def jsonld(titles_and_urls: list[tuple[object, str]]) -> bytes:
    graph: list[dict[str, object]] = [{"@id": "x", "@type": "dcat:Dataset"}]
    for title, url in titles_and_urls:
        graph.append(
            {
                "@id": f"d{len(graph)}",
                "@type": "dcat:Distribution",
                "dct:title": title,
                "dcat:accessURL": {"@id": url},
                "dct:license": {
                    "@id": "http://publications.europa.eu/resource/authority/licence/CC_BY_4_0"
                },
            }
        )
    return json.dumps({"@graph": graph, "@context": {}}).encode()


def multilingual(en: str) -> list[dict[str, str]]:
    return [
        {"@language": "fr", "@value": f"Lien vers {en}"},
        {"@language": "en", "@value": f"Link to {en}"},
    ]


def test_vol_a_resolves_the_volume_a_distribution_only() -> None:
    payload = jsonld(
        [
            (multilingual("eb_92_volume_A_AI.zip"), "https://x/ai"),
            (multilingual("eb92_volume_A.zip"), "https://x/a"),
            (multilingual("eb_92_volume_AA.zip"), "https://x/aa"),
            (multilingual("eb_92_volume_AP.zip"), "https://x/ap"),
            ("Detailed information on public opinion website", "https://x/web"),
        ]
    )
    dist = eb.vol_a(payload)
    assert dist.url == "https://x/a" and dist.filename == "eb92_volume_A.zip"
    assert dist.raw_name("STD92") == "eb_STD92_vol_a.zip"
    for title in (
        "eb_91_volume_A_xls.zip",
        "eb95_vol A.xlsx",
        "STD101_VOL_A.xlsx",
        "Eurobarometer_Standard_99_Spring 2023_volume_A.xlsx",
        "Standard Eurobarometer 104_Autumn 2025_volume A.xlsx",
    ):
        d = eb.vol_a(
            jsonld(
                [
                    (multilingual(title), "https://x/y"),
                    (multilingual("eb_vol_AA.xlsx"), "https://x/z"),
                ]
            )
        )
        assert d.url == "https://x/y"
    assert eb.vol_a(jsonld([(multilingual("STD101_VOL_A.xlsx"), "https://x/y")])).raw_name(
        "STD101"
    ) == ("eb_STD101_vol_a.xlsx")
    # a plain-string title (STD103's JSON-LD) also works
    assert eb.vol_a(jsonld([("Link to eb103_volume_A.xlsx", "https://x/s")])).url == "https://x/s"


def test_vol_a_fails_closed_when_ambiguous_missing_or_not_jsonld() -> None:
    with pytest.raises(ValueError, match="no Volume A"):
        eb.vol_a(jsonld([(multilingual("eb_volume_B.xlsx"), "https://x/b")]))
    with pytest.raises(ValueError, match="2 Volume A"):
        eb.vol_a(
            jsonld(
                [
                    (multilingual("eb_volume_A.xlsx"), "https://x/1"),
                    (multilingual("eb_vol_A.zip"), "https://x/2"),
                ]
            )
        )
    with pytest.raises(ValueError, match="JSON-LD"):
        eb.vol_a(b"<html>maintenance</html>")
    with pytest.raises(ValueError, match="JSON-LD"):
        eb.vol_a(b'{"error": "x"}')


# ---------------------------------------------------------------- workbook -> rows


def by_key(rows: list[dict[str, object]]) -> dict[tuple[object, object], dict[str, object]]:
    return {(r["iso3"], r["item"]): r for r in rows}


def test_expectation_rows_from_xlsx_take_the_four_items_and_drop_aggregates() -> None:
    rows = eb.expectation_rows(workbook(), std101())
    assert {r["item"] for r in rows} == {
        "life_general",
        "household_finance",
        "national_economy",
        "employment_situation",
    }
    iso3s = {r["iso3"] for r in rows}
    assert len(iso3s) == len(CODES) - 1 - 2 - 2  # EU27, D-W/D-E, XK/CY(TCC) dropped
    assert {"DEU", "GRC", "GBR", "MDA", "TUR", "GEO", "FIN"} <= iso3s
    assert len(rows) == 4 * len(iso3s)
    k = by_key(rows)
    fi = k[("FIN", "life_general")]
    assert fi == {
        "iso3": "FIN",
        "wave": "STD101",
        "fieldwork_year": 2024,
        "fieldwork_half": 1,
        "fieldwork_start": "2024-04",
        "item": "life_general",
        "better_share": 35.0,
        "worse_share": 12.0,
        "same_share": 50.0,
        "dk_share": 3.0,
        "source": "eurobarometer_std",
    }
    assert k[("DEU", "life_general")]["better_share"] == 17.0  # DE, not D-W / D-E
    assert k[("BEL", "life_general")]["dk_share"] == 0.0  # "-" means no respondents
    assert k[("BEL", "household_finance")]["same_share"] == 42.0  # label "Same"
    assert k[("BEL", "employment_situation")]["dk_share"] == 1.0  # label "DK"
    assert k[("BEL", "national_economy")]["better_share"] == 21.0


EB91_CODES = ["BE", "DK", "D-W", "DE", "D-E", "FI", "UK"]


def test_expectation_rows_from_legacy_xls_and_zip_wrappers(
    xls_from_sheets: Callable[[dict[str, list[list[object]]]], bytes],
) -> None:
    std91 = eb.WAVES[0]
    sheets = {
        "Index": [["Eurobarometer 91.5"]],
        "QA2a.1": [
            ["Index", "Eurobarometer 91.5"],
            ["VOLUME A Pondéré Weighted", "Terrain/Fieldwork : 07/06 - 01/07/2019"],
            ["QA2a.1 Quelles sont vos attentes ...", QUESTION_EN.replace("QA2.1.", "QA2a.1")],
            [],
            ["Votre vie en général", "Your life in general"],
            [],
            [],
            [],
            [None, None, "UE28\nEU28", "UE28-UK\nEU28-UK", *EB91_CODES],
            [None, "TOTAL", 27464, 26432, 1057, 1013, 1035, 1487, 452, 1000, 1032],
            [None, "Meilleurs", 8651, 8155, 265, 282, 235, 345, 113, 300, 400],
            [None, "Better", 0.32, 0.31, 0.25, 0.28, 0.23, 0.23, 0.25, 0.3, 0.4],
            [None, "Moins bons", 1, 1, 1, 1, 1, 1, 1, 1, 1],
            [None, "Worse", 0.1, 0.1, 0.12, 0.05, 0.06, 0.07, 0.1, 0.1, 0.1],
            [None, "Sans changement", 1, 1, 1, 1, 1, 1, 1, 1, 1],
            [None, "Same", 0.56, 0.57, 0.62, 0.66, 0.7, 0.69, 0.65, 0.58, 0.48],
            [None, "NSP", 1, 1, 1, 1, 1, 1, 1, 1, 1],
            [None, "DK", 0.02, 0.02, 0.01, 0.01, 0.01, 0.01, "- ", 0.02, 0.02],
        ],
    }
    xls = xls_from_sheets(sheets)
    rows = eb.expectation_rows(xls, std91)
    assert [(r["iso3"], r["better_share"]) for r in rows] == [
        ("BEL", 25.0),
        ("DNK", 28.0),
        ("DEU", 23.0),
        ("FIN", 30.0),
        ("GBR", 40.0),
    ]
    assert rows[0]["wave"] == "STD91" and rows[0]["fieldwork_year"] == 2019
    # the same workbook wrapped the way the portal ships it (zip with one member, or several)
    assert eb.expectation_rows(zip_of("eb_91_volume_A.xls", xls), std91) == rows
    multi = zip_of(
        "eb92_Volume_A.xls",
        xls,
        {"eb92_Volume_A_Budget.xls": b"junk", "eb92_Volume_A_with CC.xls": b"junk"},
    )
    assert eb.expectation_rows(multi, std91) == rows
    assert eb.expectation_rows(zip_of("eb_93_volume_A.xlsx", workbook()), std101()) == (
        eb.expectation_rows(workbook(), std101())
    )


def test_expectation_rows_returns_empty_when_question_absent_and_fails_on_bad_layout() -> None:
    no_question = xlsx_from_sheets(
        {"Content": [["Project:"]], "QA1_1": [["QA1.1 Trust"], [], ["x"]]}
    )
    assert eb.expectation_rows(no_question, std101()) == []
    bad_header = xlsx_from_sheets(
        {
            "QA2_1": item_sheet(
                "Your life in general",
                better=fill(0.1),
                worse=fill(0.1),
                same=fill(0.1),
                dk=fill(0.1),
            )[:8]
        }
    )
    with pytest.raises(ValueError, match="country header"):
        eb.expectation_rows(bad_header, std101())
    rows = item_sheet(
        "Your life in general", better=fill(0.1), worse=fill(0.1), same=fill(0.1), dk=fill(0.1)
    )
    rows = [r for r in rows if not r or r[0] != "Worse"]
    with pytest.raises(ValueError, match="Worse"):
        eb.expectation_rows(xlsx_from_sheets({"QA2_1": rows}), std101())
    not_a_share = item_sheet(
        "Your life in general", better=fill(26), worse=fill(0.1), same=fill(0.1), dk=fill(0.1)
    )
    with pytest.raises(ValueError, match="share"):
        eb.expectation_rows(xlsx_from_sheets({"QA2_1": not_a_share}), std101())
    with pytest.raises(ValueError, match="workbook"):
        eb.expectation_rows(b"<html>maintenance</html>", std101())
    with pytest.raises(ValueError, match="Volume A"):
        eb.expectation_rows(zip_of("readme.txt", b"hi"), std101())


# ---------------------------------------------------------------- mart


def test_build_expectations_mart_adds_net_optimism_and_sorts() -> None:
    rows = eb.expectation_rows(workbook(), std101())
    older = [
        {
            **r,
            "wave": "STD91",
            "fieldwork_year": 2019,
            "fieldwork_half": 1,
            "fieldwork_start": "2019-06",
        }
        for r in rows
    ]
    mart = eb.build_expectations_mart(rows + older)
    assert len(mart) == 2 * len(rows)
    assert set(mart[0]) == set(rows[0]) | {"net_optimism"}
    fi = [r for r in mart if r["iso3"] == "FIN" and r["item"] == "life_general"]
    assert [r["fieldwork_year"] for r in fi] == [2019, 2024]
    assert fi[1]["net_optimism"] == 23.0  # 35 - 12
    keys = [(r["iso3"], r["fieldwork_year"], r["fieldwork_half"], r["item"]) for r in mart]
    assert keys == sorted(keys)


# ---------------------------------------------------------------- layout variants in real files


def test_std93_bilingual_label_cell_with_shares_on_the_next_row() -> None:
    n = len(CODES)
    rows: list[list[object]] = [
        [None, None, None, None, None, None, None, None, "Eurobarometer - 93.1"],
        [None, None, "Volume A weighted"],
        [None, None, "QA2a.1 Quelles sont vos attentes ...", *([None] * 8), QUESTION_EN],
        [None, None, "Votre vie en général", *([None] * 8), "Your life in general"],
        ["  ", None, "Base: Ensemble"],
        [],
        [],
        [],
        ["<<Back to content", None, "UE27 EU27", *CODES[1:]],
        [" ", "Total", 26681, *([1000] * (n - 1))],
        [],
        [None, "Meilleurs\nBetter", 6570, *([218] * (n - 1))],
        [None, None, 0.25, *([0.22] * (n - 1))],
        [None, "Moins bons\nWorse", 4029, *([175] * (n - 1))],
        [None, None, 0.15, *([0.17] * (n - 1))],
        [None, "Sans changement\nSame", 15321, *([603] * (n - 1))],
        [None, None, 0.57, *([0.6] * (n - 1))],
        [None, "Ne sait pas\nDon't know", 761, *([12] * (n - 1))],
        [None, None, 0.03, *([0.01] * (n - 1))],
    ]
    std93 = next(w for w in eb.WAVES if w.code == "STD93")
    got = eb.expectation_rows(xlsx_from_sheets({"T17": rows}), std93)
    assert len(got) == len(CODES) - 1 - 2 - 2
    be = next(r for r in got if r["iso3"] == "BEL")
    assert (be["better_share"], be["worse_share"], be["same_share"], be["dk_share"]) == (
        22.0, 17.0, 60.0, 1.0,
    )  # fmt: skip


def test_cy_tcc_only_sheets_and_the_std100_economy_wording() -> None:
    tcc: list[list[object]] = [
        ["Index", *([None] * 5), "Eurobarometer 91.5"],
        [],
        [None, "QA2b.1 Quelles ...", *([None] * 9), QUESTION_EN],
        [],
        [None, "Votre vie en général", *([None] * 9), "Your life in general"],
        [],
        [],
        [],
        [None, None, None, "CY (tcc)"],
        [None, "TOTAL", None, 500],
        [None, "Meilleurs", None, 190],
        [None, "Better", None, 0.38],
        [None, "Moins bons", None, 120],
        [None, "Worse", None, 0.24],
        [None, "Sans changement", None, 146],
        [None, "Same", None, 0.29],
        [None, "NSP", None, 44],
        [None, "DK", None, 0.09],
    ]
    economy = item_sheet(
        "The state of (OUR COUNTRY)'s economy  ",
        better=fill(0.11), worse=fill(0.41), same=fill(0.44), dk=fill(0.04),
        question="QA2.3. What are your expectations for the next 12 months: will they be ...?",
    )  # fmt: skip
    got = eb.expectation_rows(xlsx_from_sheets({"QA2b.1": tcc, "QA2_3": economy}), std101())
    assert {r["item"] for r in got} == {"national_economy"}
    assert next(r for r in got if r["iso3"] == "FIN")["worse_share"] == 41.0
